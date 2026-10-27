"""Rebuilds the data block in APEX.html from the Excel files in ./data.

All paths are relative to this script's folder:
  data/DI Ageing-Pends.xlsx  -> table "Table2" (sheet "Update Here Please")
  data/RPTDB.xlsx            -> table "RPT"    (sheet "rpt")
  APEX.html                  -> block between <!--APEX_DATA--> and <!--/APEX_DATA-->
"""
import json
import os
import re
import sys
import warnings
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

warnings.filterwarnings('ignore', module='openpyxl')

BASE = Path(__file__).resolve().parent
DI_FILE = BASE / 'data' / 'DI Ageing-Pends.xlsx'
RPT_FILE = BASE / 'data' / 'RPTDB.xlsx'
HTML_FILE = BASE / 'APEX.html'

MANAGER_ALIASES = {'CI-79131274': 'Clarissa Alfonso', 'Clarisa Alfonso': 'Clarissa Alfonso'}


def clean(v):
    return v.strip() if isinstance(v, str) else ('' if v is None else v)


def norm_mgr(v):
    v = clean(v)
    return MANAGER_ALIASES.get(v, v)


def table_rows(path, table_name, columns):
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        if table_name in ws.tables:
            cells = ws[ws.tables[table_name].ref]
            header = [str(clean(c.value)) for c in cells[0]]
            missing = [c for c in columns if c not in header]
            if missing:
                sys.exit(f'ERROR: {path.name} / {table_name} is missing columns: {missing}')
            idx = [header.index(c) for c in columns]
            return [tuple(clean(r[i].value) for i in idx) for r in cells[1:]]
    sys.exit(f'ERROR: table "{table_name}" not found in {path.name}')


def build_di():
    rows = table_rows(DI_FILE, 'Table2', ['Status', 'Category', 'Manager', 'Sub-status'])
    cat_stat = {c: Counter() for c in ('Pends', 'Ageing')}
    mgr_total, mgr_pends, mgr_age = Counter(), Counter(), Counter()
    matrix, sub_total = Counter(), Counter()

    for status, cat, mgr, sub in rows:
        mgr = norm_mgr(mgr) or '<blank>'
        sub = sub or '<blank>'
        if cat:
            key = 'completed' if status == 'Completed' else 'pending' if status == 'Pending' else 'blank'
            cat_stat.setdefault(cat, Counter())[key] += 1
        mgr_total[mgr] += 1
        if cat == 'Pends':
            mgr_pends[mgr] += 1
        if cat == 'Ageing':
            mgr_age[mgr] += 1
        matrix[(sub, mgr)] += 1
        sub_total[sub] += 1

    managers = sorted((m for m in mgr_total if m != '<blank>'), key=lambda m: -mgr_total[m])

    def card(cat):
        s = cat_stat.get(cat, Counter())
        return {'total': sum(s.values()), 'completed': s['completed'], 'pending': s['pending'], 'blank': s['blank']}

    matrix_rows = [{'s': s, 'cells': [matrix[(s, m)] for m in managers], 'total': sub_total[s]}
                   for s in sorted(sub_total, key=lambda k: -sub_total[k])]
    data = {
        'pends': card('Pends'),
        'ageing': card('Ageing'),
        'managers': managers,
        'mgrPends': [{'m': m, 'v': mgr_pends[m]} for m in managers],
        'mgrAge': [{'m': m, 'v': mgr_age[m]} for m in managers],
        'matrix': {'rows': matrix_rows, 'colTotals': [mgr_total[m] for m in managers], 'grand': sum(sub_total.values())},
    }

    # sanity checks
    assert data['matrix']['grand'] == len(rows)
    assert all(sum(r['cells']) + matrix[(r['s'], '<blank>')] == r['total'] for r in matrix_rows)
    assert sum(data['matrix']['colTotals']) + mgr_total['<blank>'] == len(rows)
    for c in ('pends', 'ageing'):
        d = data[c]
        assert d['completed'] + d['pending'] + d['blank'] == d['total']
    return data, len(rows)


def build_review():
    rows = table_rows(RPT_FILE, 'RPT', ['AM', 'Pend Review'])
    per = defaultdict(Counter)
    for am, review in rows:
        am = norm_mgr(am)
        if not am:
            continue
        per[am]['v'] += review == 'Valid'
        per[am]['i'] += review == 'Invalid'
    out = [{'m': m, 'v': c['v'], 'i': c['i'], 't': c['v'] + c['i']}
           for m, c in sorted(per.items(), key=lambda kv: -(kv[1]['v'] + kv[1]['i']))]
    totals = {k: sum(r[k] for r in out) for k in ('v', 'i', 't')}
    assert totals['v'] + totals['i'] == totals['t']
    return {'rows': out, 'totals': totals}, len(rows)


def mtime(p):
    return datetime.fromtimestamp(p.stat().st_mtime).strftime('%Y-%m-%d %H:%M')


def main():
    for p in (DI_FILE, RPT_FILE, HTML_FILE):
        if not p.exists():
            sys.exit(f'ERROR: not found: {p.relative_to(BASE)}')

    di, di_count = build_di()
    review, rpt_count = build_review()
    payload = {
        'generated': datetime.now().strftime('%Y-%m-%d %H:%M'),
        'sources': [
            {'name': DI_FILE.name, 'modified': mtime(DI_FILE), 'rows': di_count},
            {'name': RPT_FILE.name, 'modified': mtime(RPT_FILE), 'rows': rpt_count},
        ],
        **di,
        'review': review,
    }

    blob = json.dumps(payload, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    html = HTML_FILE.read_text(encoding='utf-8')
    pattern = re.compile(r'(<!--APEX_DATA-->).*?(<!--/APEX_DATA-->)', re.S)
    if not pattern.search(html):
        sys.exit('ERROR: <!--APEX_DATA--> marker not found in APEX.html')
    html = pattern.sub(lambda m: f'{m.group(1)}\n<script>var APEX_DATA = {blob};</script>\n{m.group(2)}', html)

    tmp = HTML_FILE.with_suffix('.html.tmp')
    tmp.write_text(html, encoding='utf-8')
    os.replace(tmp, HTML_FILE)

    p, a, r = payload['pends'], payload['ageing'], review['totals']
    print(f"APEX.html updated {payload['generated']}")
    print(f"  Pends  total={p['total']} completed={p['completed']} pending={p['pending']} blank={p['blank']}")
    print(f"  Ageing total={a['total']} completed={a['completed']} pending={a['pending']} blank={a['blank']}")
    print(f"  Substatus grand total={payload['matrix']['grand']}  managers={payload['managers']}")
    print(f"  Pend Reviewed valid={r['v']} invalid={r['i']} total={r['t']}")


if __name__ == '__main__':
    main()
